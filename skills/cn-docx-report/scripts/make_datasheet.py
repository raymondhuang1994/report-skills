# -*- coding: utf-8 -*-
"""make_datasheet.py —— 生成 Excel 取数单（宿主没有 xlsx 技能时的兜底）。
用法：python3 make_datasheet.py 取数需求.json 输出.xlsx
JSON：{"items": [{"编号": "G1", "指标": "...", ...}, ...], "review": [{"编号": "R1", "数字": "...", "出处": "...", "所在位置": "..."}]}
列定义见 references/evidence-discipline.md 第 8 节；缺的列留空，不要删列。"""
import sys, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

COLS = ["编号", "指标", "口径定义", "标的与代码", "期间与频率", "首选来源", "备选来源", "用途", "优先级", "状态", "回填位置"]
REVIEW = ["编号", "数字", "出处", "所在位置", "复核结果"]

def sheet(ws, cols, rows, widths):
    ws.append(cols)
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="12355B"); c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows: ws.append([str(r.get(c, "")) for c in cols])
    for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"

def main():
    if len(sys.argv) < 3: raise SystemExit(__doc__)
    spec = json.load(open(sys.argv[1], encoding="utf-8")); wb = Workbook()
    ws = wb.active; ws.title = "取数单"; items = spec.get("items", [])
    sheet(ws, COLS, items, [8, 22, 28, 22, 18, 14, 14, 30, 8, 10, 14])
    n = max(len(items) + 1, 2)
    for col, opts in (("I", '"高,中,低"'), ("J", '"待取,已取,无法取得"')):
        dv = DataValidation(type="list", formula1=opts, allow_blank=True); ws.add_data_validation(dv); dv.add(f"{col}2:{col}{n + 200}")
    sheet(wb.create_sheet("复核清单"), REVIEW, spec.get("review", []), [8, 30, 40, 20, 16])
    wb.save(sys.argv[2]); print("written", sys.argv[2], f"取数 {len(items)} 行，复核 {len(spec.get('review', []))} 行")

if __name__ == "__main__": main()
