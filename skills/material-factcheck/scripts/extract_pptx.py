# -*- coding: utf-8 -*-
"""extract_pptx.py —— 把 PPT 每页的标题、文字、表格和图表数据抽成 JSON，供材料核查与复算。
用法：python3 extract_pptx.py deck.pptx [out.json] [--slides 5,29]
图表数据（饼图、柱图的类别与数值）只存在于图表对象里，文字提取工具拿不到，核对“饼图 vs 表格加总”一类口径矛盾必须用它。
依赖：python-pptx（pip install python-pptx）。"""
import sys, json
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

def walk(shapes):  # 递归展开组合形状
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP: yield from walk(sh.shapes)
        else: yield sh

def main():
    if len(sys.argv) < 2: raise SystemExit(__doc__)
    args = sys.argv[1:]; only = None
    if "--slides" in args:
        i = args.index("--slides"); only = {int(x) for x in args[i + 1].split(",")}; del args[i:i + 2]
    prs = Presentation(args[0]); out = []
    for n, slide in enumerate(prs.slides, 1):
        if only and n not in only: continue
        rec = {"slide": n, "title": "", "texts": [], "tables": [], "charts": []}
        try: rec["title"] = (slide.shapes.title.text_frame.text.strip() if slide.shapes.title else "")
        except Exception: pass
        for sh in walk(slide.shapes):
            if sh.has_text_frame and sh.text_frame.text.strip() and sh.text_frame.text.strip() != rec["title"]:
                rec["texts"].append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                rec["tables"].append([[c.text.strip() for c in r.cells] for r in sh.table.rows])
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                for plot in ch.plots:
                    cats = [str(c) for c in plot.categories]
                    series = [{"name": s.name, "values": [None if v is None else round(float(v), 4) for v in s.values]} for s in plot.series]
                    rec["charts"].append({"chart_type": str(ch.chart_type), "categories": cats, "series": series})
        out.append(rec)
        print(f"第 {n:2d} 页  文字 {len(rec['texts']):2d}  表格 {len(rec['tables'])}  图表 {len(rec['charts'])}  {rec['title'][:40]}")
    if len(args) > 1:
        json.dump(out, open(args[1], "w", encoding="utf-8"), ensure_ascii=False, indent=1); print("written", args[1])

if __name__ == "__main__": main()
